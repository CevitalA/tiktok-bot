import openpyxl
import json
import re
import os
import io
import base64
from datetime import datetime

ALL_LEAGUES = [
    "A1", "A2", "A3", "A4", "A5",
    "B1", "B2", "B3", "B4", "B5",
    "C1", "C2", "C3", "C4", "C5",
    "D1", "D2", "D3", "D4", "D5"
]

OUTPUT_DATA_FILE = os.path.join(os.path.dirname(__file__), "data", "leagues_data.json")

def normalize_text(text: str) -> str:
    if text is None:
        return ""
    return str(text).strip().lower().replace(" ", "").replace("_", "").replace("-", "")

def reset_leagues_data():
    os.makedirs(os.path.dirname(OUTPUT_DATA_FILE), exist_ok=True)
    clean_payload = {
        "updated_at": None,
        "total_accounts": 0,
        "total_earnings": 0.0,
        "total_diamonds": 0,
        "active_leagues": [],
        "leagues": {lg: [] for lg in ALL_LEAGUES}
    }
    with open(OUTPUT_DATA_FILE, "w", encoding="utf-8") as f:
        json.dump(clean_payload, f, ensure_ascii=False, indent=2)
    return clean_payload

def parse_excel_file(file_path: str) -> dict:
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"File not found: {file_path}")
        
    wb = openpyxl.load_workbook(file_path, data_only=True)
    
    target_sheet = None
    for name in wb.sheetnames:
        norm_name = name.lower()
        if "rank" in norm_name or "all" in norm_name or "ليج" in norm_name or "ترتيب" in norm_name:
            target_sheet = wb[name]
            break
    if not target_sheet:
        target_sheet = wb.active

    # التحقق من وجود صور مدمجة داخل الشيت (Embedded Images)
    embedded_images_by_row = {}
    if hasattr(target_sheet, "_images"):
        for img in target_sheet._images:
            try:
                row_idx = img.anchor._from.row + 1
                img_data = img.ref.getvalue()
                img_format = getattr(img, "format", "png")
                b64 = f"data:image/{img_format};base64," + base64.b64encode(img_data).decode("ascii")
                embedded_images_by_row[row_idx] = b64
            except Exception:
                pass

    header_map = {}
    # Use values_only=False so we can extract cell.hyperlink if present
    rows_iter = target_sheet.iter_rows(values_only=False)
    
    first_row_cells = next(rows_iter, None)
    if not first_row_cells:
        raise ValueError("ملف الإكسل فارغ تماماً!")
        
    for idx, cell in enumerate(first_row_cells):
        col_norm = normalize_text(cell.value)
        header_map[idx] = col_norm

    col_mapping = {
        "rank": ["rank", "الترتيب", "ترتيب", "المركز", "#"],
        "creatorName": ["creatorname", "creator", "name", "الاسم", "اسم", "اسمالمنشئ", "منشئ"],
        "username": ["username", "user", "اليوزر", "اسمالمستخدم", "يوزر", "حساب"],
        "league": ["league", "الليج", "الدوري", "المستوى", "ليج", "الدرجة"],
        "earnings": ["estearnings($)", "estearnings", "earnings", "earning", "ارباح", "الارباح", "الدولار", "usd", "$", "profit"],
        "diamonds": ["diamonds", "diamond", "الماس", "الالماس", "الألماس"],
        "eligibility": ["eligibility", "حالة", "الحالة", "الاهلية", "eligible"],
        "creatorUrl": [
            "creatorurl", "creatorlink", "creatoruri", "profileurl", "profilelink",
            "tiktokurl", "tiktoklink", "userurl", "accounturl", "رابطالمنشئ", "لينكالمنشئ",
            "رابطالحساب", "لينكالحساب", "رابطالتيكتوك", "لينكالتيكتوك", "رابطالصفحة",
            "رابط", "لينك", "url", "link", "tiktok"
        ],
        "avatar": [
            "avatar", "avatarurl", "image", "photo", "picture", "pic", "profileimage",
            "profilepic", "useravatar", "صورة", "الصورة", "صورةالمنشئ", "صورةالمستخدم",
            "صورةالبروفايل", "رابطالصورة", "افاتار", "أفاتار", "لوجو"
        ]
    }

    def find_col_idx(field_name):
        candidates = col_mapping[field_name]
        # Priority 1: exact normalized match
        for idx, col_str in header_map.items():
            if col_str in candidates:
                return idx
        # Priority 2: candidate is contained in col_str or vice versa (with disambiguation)
        for idx, col_str in header_map.items():
            # Disambiguation: avoid creatorUrl matching image/avatar columns
            if field_name == "creatorUrl" and any(img_w in col_str for img_w in ["avatar", "image", "photo", "picture", "pic", "صورة", "افاتار", "لوجو"]):
                continue
            if field_name == "avatar" and "creator" in col_str and "url" in col_str and "avatar" not in col_str:
                continue
            for cand in candidates:
                if cand == col_str or cand in col_str or col_str in cand:
                    return idx
        return None

    resolved_cols = {f: find_col_idx(f) for f in col_mapping}
    
    if resolved_cols["league"] is None:
        raise ValueError("لم يتم العثور على عمود الليج (League) في الملف!")

    leagues_dict = {lg: [] for lg in ALL_LEAGUES}
    total_processed = 0
    total_earnings = 0.0
    total_diamonds = 0

    for r_idx, row_cells in enumerate(rows_iter, start=2):
        if not row_cells or not any(c.value is not None for c in row_cells):
            continue
            
        def get_val(f_name):
            c_idx = resolved_cols[f_name]
            if c_idx is not None and c_idx < len(row_cells):
                cell = row_cells[c_idx]
                if f_name == "creatorUrl":
                    # Check for Excel hyperlink target first
                    if cell.hyperlink and cell.hyperlink.target:
                        return cell.hyperlink.target
                return cell.value
            return None

        league_raw = get_val("league")
        if not league_raw:
            continue
            
        league = str(league_raw).strip().upper()
        if league not in ALL_LEAGUES:
            continue

        raw_name = get_val("creatorName")
        raw_user = get_val("username")
        raw_rank = get_val("rank")
        raw_earn = get_val("earnings")
        raw_diam = get_val("diamonds")
        raw_elig = get_val("eligibility")
        raw_url = get_val("creatorUrl")
        raw_avatar = get_val("avatar")

        username = str(raw_user or f"user_{r_idx}").strip().lstrip("@")
        creator_name = str(raw_name or username).strip()
        
        # Clean and format Creator URL
        clean_url = str(raw_url).strip() if raw_url is not None else ""
        clean_url = clean_url.strip("'\" \t\r\n\u200e\u200f")
        if clean_url and clean_url.lower() not in ["none", "null", "nan", "-", ""]:
            if clean_url.lower().startswith("http://") or clean_url.lower().startswith("https://"):
                creator_url = clean_url
            elif "tiktok.com" in clean_url.lower():
                creator_url = "https://" + clean_url.lstrip("/").lstrip(":")
            elif clean_url.startswith("@"):
                creator_url = f"https://www.tiktok.com/{clean_url}"
            elif "." in clean_url and "/" in clean_url:
                creator_url = "https://" + clean_url
            else:
                creator_url = f"https://www.tiktok.com/@{clean_url.lstrip('@')}"
        else:
            creator_url = f"https://www.tiktok.com/@{username}"

        # استخراج الصورة إذا كانت موجودة في العمود أو مدمجة بالشيت
        avatar_url = ""
        if raw_avatar:
            s_av = str(raw_avatar).strip()
            if s_av.startswith("http") or s_av.startswith("data:image") or s_av.startswith("/"):
                avatar_url = s_av
        if not avatar_url and r_idx in embedded_images_by_row:
            avatar_url = embedded_images_by_row[r_idx]

        try:
            rank = int(float(str(raw_rank).replace(",", "")))
        except:
            rank = len(leagues_dict[league]) + 1

        try:
            earnings = float(re.sub(r"[^0-9.-]", "", str(raw_earn or 0)))
        except:
            earnings = 0.0

        try:
            diamonds = int(float(re.sub(r"[^0-9.-]", "", str(raw_diam or 0))))
        except:
            diamonds = 0

        eligibility = str(raw_elig or "Available").strip()

        item = {
            "rank": rank,
            "creatorName": creator_name,
            "username": username,
            "league": league,
            "earnings": earnings,
            "diamonds": diamonds,
            "eligibility": eligibility,
            "creatorUrl": creator_url,
            "avatar": avatar_url
        }

        leagues_dict[league].append(item)
        total_processed += 1
        total_earnings += earnings
        total_diamonds += diamonds

    for lg in ALL_LEAGUES:
        leagues_dict[lg].sort(key=lambda x: x["rank"])

    os.makedirs(os.path.dirname(OUTPUT_DATA_FILE), exist_ok=True)
    payload = {
        "updated_at": datetime.now().isoformat(),
        "total_accounts": total_processed,
        "total_earnings": round(total_earnings, 2),
        "total_diamonds": total_diamonds,
        "active_leagues": [lg for lg in ALL_LEAGUES if len(leagues_dict[lg]) > 0],
        "leagues": leagues_dict
    }

    with open(OUTPUT_DATA_FILE, "w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False, indent=2)

    return payload

def get_current_leagues_data() -> dict:
    if os.path.exists(OUTPUT_DATA_FILE):
        try:
            with open(OUTPUT_DATA_FILE, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            pass
    return reset_leagues_data()
